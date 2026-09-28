"""Build the Word report from outputs/results.json and verified figures."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "outputs" / "results.json"
FIG = ROOT / "outputs" / "figures"
DEST = ROOT / "Лабораторная_работа_1_газовая_турбина.docx"
DETAILS = ROOT / "report_details.json"
BLACK = RGBColor(0, 0, 0)


def num(value, digits=3):
    return f"{value:.{digits}f}".replace(".", ",")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def set_cell_margin(cell, top=85, start=90, bottom=85, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def format_cell(cell, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            run.font.bold = bold
            run.font.color.rgb = BLACK


def table(doc, caption, headers, rows, widths=None, center_cols=()):
    cap = doc.add_paragraph(caption)
    cap.style = "Table Caption"
    cap.paragraph_format.keep_with_next = True
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    header_pr = tbl.rows[0]._tr.get_or_add_trPr()
    header_pr.append(OxmlElement("w:tblHeader"))
    for i, name in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = name
        shade(cell, "E7E7E7")
        format_cell(cell, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
        if widths:
            cell.width = Cm(widths[i])
    for values in rows:
        row = tbl.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            format_cell(cell, align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
            if widths:
                cell.width = Cm(widths[i])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.5
    return tbl


def figure(doc, filename, caption, width=16.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(FIG / filename), width=Cm(width))
    cap = doc.add_paragraph(caption)
    cap.style = "Caption"
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_together = True


def setup(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    cap = doc.styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(12)
    cap.font.color.rgb = BLACK
    cap.font.italic = False
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = 1.0
    from docx.enum.style import WD_STYLE_TYPE
    tabcap = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    tabcap.base_style = normal
    tabcap.font.name = "Times New Roman"
    tabcap.font.size = Pt(12)
    tabcap.font.color.rgb = BLACK
    tabcap.paragraph_format.first_line_indent = Cm(0)
    tabcap.paragraph_format.space_before = Pt(8)
    tabcap.paragraph_format.space_after = Pt(4)
    tabcap.paragraph_format.line_spacing = 1.0
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    run = footer.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = BLACK
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def p(doc, text, bold_start=None):
    para = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        para.add_run(bold_start).bold = True
        para.add_run(text[len(bold_start):])
    else:
        para.add_run(text)
    return para


def list_item(doc, text):
    """Create a portable list item without Word's font-dependent bullet glyph."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(1.25)
    para.paragraph_format.first_line_indent = Cm(-0.75)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.5
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run("— " + text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = BLACK
    return para


def title_page(doc, details):
    def line(text, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0):
        paragraph = doc.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = BLACK
        return paragraph

    line("Министерство науки и высшего образования Российской Федерации", size=12)
    line(details["university"], size=14, bold=True, before=10)
    line(details["faculty"], size=14, before=6)
    line("ОТЧЁТ", size=16, bold=True, before=135)
    line("по лабораторной работе № 1", size=14, before=8)
    line("по дисциплине «Математические методы теории систем»", size=14, before=4)
    line("Линейная многомерная регрессия: прогноз выбросов CO газовой турбины", size=14, bold=True, before=10)
    line(f"Выполнила: студентка группы {details['group']}", align=WD_ALIGN_PARAGRAPH.RIGHT, before=105)
    line(details["student"], align=WD_ALIGN_PARAGRAPH.RIGHT, before=4)
    line(f"Проверил: {details['teacher']}", align=WD_ALIGN_PARAGRAPH.RIGHT, before=15)
    line(f"{details['city']} — {details['year']}", before=95)
    doc.add_page_break()


def section_heading(doc, number, title, first=False):
    if not first:
        doc.add_page_break()
    doc.add_heading(f"{number} {title}", level=1)


def main():
    if not RESULTS.exists():
        raise FileNotFoundError("Run analysis.py before make_report.py")
    if not DETAILS.exists():
        raise FileNotFoundError("Missing project file: report_details.json")
    r = json.loads(RESULTS.read_text(encoding="utf-8"))
    details = json.loads(DETAILS.read_text(encoding="utf-8"))
    q = r["data_quality"]
    hold = r["holdout"]
    diag = r["gauss_markov_diagnostics_for_full_ols"]
    candidates = {x["model"]: x for x in r["candidate_models"]}
    doc = Document()
    setup(doc)
    title_page(doc, details)
    best = hold["selected_overall"]

    section_heading(doc, 1, "Постановка задачи и источник данных", first=True)
    p(
        doc,
        f"На 36 733 почасовых измерениях газовой турбины построена модель концентрации "
        f"угарного газа (CO) по девяти параметрам среды и установки. По последовательной проверке "
        f"на 2012 и 2013 годах выбрана Ridge-регрессия с квадратичными признаками. На отложенных "
        f"2014–2015 годах: MAE = {num(best['MAE_mg_m3'])} мг/м³, "
        f"RMSE = {num(best['RMSE_mg_m3'])} мг/м³, R² = {num(best['R2'])}. "
        "Проверки предпосылок показали сильную связь части датчиков, неодинаковую дисперсию и "
        "временную зависимость ошибок обычной линейной регрессии."
    )
    p(
        doc,
        "Цель — оценить концентрацию CO в выхлопе турбины (мг/м³) по параметрам окружающей "
        "среды и технологического режима. Практический сценарий — дополнительный сигнал для "
        "анализа выбросов и обнаружения нетипичной работы установки. Целевая величина непрерывная; "
        "поэтому задача относится к регрессии. Данные одной установки не позволяют заявить, что "
        "модель подходит для других турбин без отдельной проверки."
    )
    p(
        doc,
        "Источник данных [1]: Gas Turbine CO and NOx Emission Data Set, UCI Machine Learning Repository, "
        "DOI 10.24432/C5WC95. Наблюдения получены на газовой турбине в Турции в 2011–2015 годах. "
        "По описанию UCI значения датчиков агрегированы за час; внутри файлов строки упорядочены "
        "по времени, но точных временных меток нет. Лицензия набора данных — CC BY 4.0."
    )

    section_heading(doc, 2, "Структура данных и подготовка")
    p(
        doc,
        "Пять исходных файлов gt_2011.csv … gt_2015.csv содержат по 11 числовых столбцов. "
        "В коде из имени файла добавлен столбец year для разбиения по годам. В общей таблице "
        f"{q['rows']:,} строк и {q['total_missing_values']} пропусков. Найдено "
        f"{q['exact_duplicate_measurements_with_year']} полностью одинаковых строк измерений "
        "(с учётом года). Они оставлены: без точных временных меток нельзя установить, являются ли "
        "они техническими дублями или повторившимся состоянием установки."
    )
    table(
        doc,
        "Таблица 1 — Число наблюдений и статистики CO по годам",
        ["Год", "Число строк", "Средний CO, мг/м³", "Медиана CO, мг/м³"],
        [
            (year, q["rows_by_year"][year]["rows"], num(q["rows_by_year"][year]["CO_mean_mg_m3"]), num(q["rows_by_year"][year]["CO_median_mg_m3"]))
            for year in map(str, range(2011, 2016))
        ],
        widths=[2, 3.5, 5.5, 5.5],
        center_cols=(0, 1, 2, 3),
    )
    p(
        doc,
        f"Распределение CO скошено вправо: медиана {num(q['CO_median_mg_m3'])} мг/м³, "
        f"99-й процентиль {num(q['CO_p99_mg_m3'])} мг/м³, максимум {num(q['CO_max_mg_m3'])} мг/м³. "
        "Редкие высокие значения важны для интерпретации RMSE и графика ошибок."
    )
    figure(doc, "target_distribution.png", "Рисунок 1 — Распределение CO и различия между годами")

    doc.add_heading("2.1 Признаки и техническое описание", level=2)
    table(
        doc,
        "Таблица 2 — Поля исходного набора данных, использованные в модели",
        ["Столбец", "Содержание", "Единица"],
        [
            ("AT", "Температура окружающего воздуха", "°C"),
            ("AP", "Атмосферное давление", "мбар"),
            ("AH", "Влажность окружающего воздуха", "%"),
            ("AFDP", "Перепад давления на воздушном фильтре", "мбар"),
            ("GTEP", "Давление выхлопа газовой турбины", "мбар"),
            ("TIT", "Температура на входе турбины", "°C"),
            ("TAT", "Температура после турбины", "°C"),
            ("TEY", "Энергетическая выработка турбины", "МВт·ч"),
            ("CDP", "Давление на выходе компрессора", "мбар"),
            ("CO", "Концентрация угарного газа — цель", "мг/м³"),
        ],
        widths=[2.7, 10.8, 3],
        center_cols=(0, 2),
    )
    p(
        doc,
        "NOX — вторая концентрация выбросов в исходном файле. Она не включена в признаки: "
        "постановка предполагает оценку CO без обращения к другому датчику выбросов. "
        "Исходные CSV читаются через pandas; весь процесс обработки и оценки описан в analysis.py."
    )

    section_heading(doc, 3, "План эксперимента и испытанные методы")
    p(
        doc,
        "Временное разбиение предотвращает попадание будущих режимов работы в обучение. Первый "
        "шаг проверки: обучение на 2011 году, оценка на 2012-м. Второй: обучение на 2011–2012 "
        "годах, оценка на 2013-м. Критерий выбора — среднее MAE этих двух шагов. Затем модель "
        "обучена заново на 2011–2013 годах и один раз проверена на 2014–2015 годах. "
        "Стандартизация признаков каждый раз оценивается только на обучающей части."
    )
    labels = {
        "mean_baseline": "Среднее обучающей выборки",
        "ambient_ols": "OLS: только среда",
        "full_ols": "OLS: все 9 признаков",
        "log_target_ols": "OLS: log(1 + CO)",
        "full_ridge_100": "Ridge: 9 признаков, α=100",
        "quadratic_ols": "OLS: квадратичные",
    }
    for alpha in (1, 10, 100, 1000, 10000):
        labels[f"quadratic_ridge_{alpha}"] = f"Ridge: квадратичные, α={alpha}"
    ordered = ["mean_baseline", "ambient_ols", "full_ols", "log_target_ols", "full_ridge_100", "quadratic_ols"] + [f"quadratic_ridge_{a}" for a in (1, 10, 100, 1000, 10000)]
    table(
        doc,
        "Таблица 3 — Результаты последовательной проверки моделей",
        ["Модель", "MAE 2012", "MAE 2013", "Среднее MAE"],
        [
            (
                labels[name],
                num(candidates[name]["folds"][0]["MAE_mg_m3"]),
                num(candidates[name]["folds"][1]["MAE_mg_m3"]),
                num(candidates[name]["mean_validation_MAE_mg_m3"]),
            )
            for name in ordered
        ],
        widths=[8.7, 2.5, 2.5, 2.8],
        center_cols=(1, 2, 3),
    )
    p(
        doc,
        "Все значения MAE в таблице 3 выражены в мг/м³. Наименьшее среднее MAE — у "
        f"квадратичной Ridge-модели с α=1000: {num(candidates['quadratic_ridge_1000']['mean_validation_MAE_mg_m3'])}. "
        "Близкий результат α=10000 показывает, что выбор коэффициента регуляризации не является "
        "абсолютно однозначным. Простая OLS и вариант с логарифмом цели использованы как "
        "контрольные решения; квадратичная OLS без регуляризации оказалась менее устойчивой."
    )

    section_heading(doc, 4, "Проверка предпосылок Гаусса—Маркова")
    p(
        doc,
        "Предпосылки проверены для обычной OLS-модели со всеми девятью исходными признаками, "
        "обученной на 2011–2013 годах. Они относятся к свойствам оценок OLS; итоговая Ridge-модель "
        "намеренно вводит смещение коэффициентов ради устойчивости прогноза."
    )
    figure(doc, "vif.png", "Рисунок 2 — VIF исходных признаков в модели OLS", width=13)
    table(
        doc,
        "Таблица 4 — Диагностика предпосылок для модели OLS",
        ["Предпосылка", "Проверка и результат", "Вывод"],
        [
            ("Линейность по коэффициентам", "OLS линейна по β; сравнение с квадратичными термами", "Простая форма ограничивает качество"),
            ("Полный ранг X", f"Ранг {diag['matrix_rank']} из {diag['parameter_count']}", "Точной линейной зависимости нет"),
            ("Практическая мультиколлинеарность", f"Максимальный VIF = {num(max(diag['vif'].values()), 1)} (TEY)", "Коэффициенты OLS нестабильны"),
            ("Постоянная дисперсия ошибок", f"Бройш–Паган: LM = {num(diag['breusch_pagan_LM'], 1)}; p < 10⁻¹⁴⁰", "Гомоскедастичность отвергается"),
            ("Независимость ошибок", f"Дарбин–Уотсон = {num(diag['durbin_watson'], 2)}", "Есть признаки положительной автокорреляции"),
            ("Нулевая условная средняя", "По этим данным не проверяется напрямую", "Возможны пропущенные факторы"),
        ],
        widths=[5, 6, 5.5],
    )
    p(
        doc,
        "Коэффициент Дарбина–Уотсона рассчитан по порядку строк, так как точных временных меток нет. "
        "Нормальность ошибок не входит в условия теоремы Гаусса—Маркова; критерий Жарка—Бера "
        "отвергает её (p практически равно нулю). Пропущенные факторы исключить нельзя."
    )
    section_heading(doc, 5, "Итоговая модель и оценка качества")
    p(
        doc,
        "Итоговая модель — Ridge-регрессия с девятью исходными признаками, их квадратами и "
        "попарными произведениями: всего 54 слагаемых. Сначала каждый исходный признак "
        "стандартизуется; затем строятся квадратичные слагаемые и они снова стандартизуются. "
        "Прогноз равен свободному члену плюс сумма произведений каждого из 54 подготовленных "
        "признаков на свой коэффициент. "
        "Коэффициенты оцениваются с L2-штрафом α = 1000. Модель линейна по коэффициентам, "
        "хотя зависимость от исходных показаний датчиков может быть криволинейной."
    )
    figure(doc, "test_fit_residuals.png", "Рисунок 3 — Фактический CO, прогноз и ошибки на 2014–2015 годах", width=13)
    p(
        doc,
        "Подбор числа признаков, вида преобразования и α выполнен только на проверочных годах "
        "2012–2013. Под «калибровкой» здесь понимается окончательная оценка параметров "
        "выбранной структуры на всех данных 2011–2013 годов. Полный набор коэффициентов "
        "записывается кодом в outputs/model_coefficients.csv; из-за сильной связи датчиков "
        "их нельзя трактовать как независимые причинные эффекты."
    )
    table(
        doc,
        "Таблица 5 — Качество прогноза на отложенных данных",
        ["Модель / период", "MAE", "RMSE", "R²", "Сдвиг прогноза"],
        [
            ("Итоговая, 2014–2015", num(best["MAE_mg_m3"]), num(best["RMSE_mg_m3"]), num(best["R2"]), num(best["bias_pred_minus_actual_mg_m3"])),
            ("Итоговая, 2014", num(hold["selected_2014"]["MAE_mg_m3"]), num(hold["selected_2014"]["RMSE_mg_m3"]), num(hold["selected_2014"]["R2"]), num(hold["selected_2014"]["bias_pred_minus_actual_mg_m3"])),
            ("Итоговая, 2015", num(hold["selected_2015"]["MAE_mg_m3"]), num(hold["selected_2015"]["RMSE_mg_m3"]), num(hold["selected_2015"]["R2"]), num(hold["selected_2015"]["bias_pred_minus_actual_mg_m3"])),
            ("OLS, 2014–2015", num(hold["full_ols_overall"]["MAE_mg_m3"]), num(hold["full_ols_overall"]["RMSE_mg_m3"]), num(hold["full_ols_overall"]["R2"]), num(hold["full_ols_overall"]["bias_pred_minus_actual_mg_m3"])),
            ("Среднее, 2014–2015", num(hold["training_mean_baseline_overall"]["MAE_mg_m3"]), num(hold["training_mean_baseline_overall"]["RMSE_mg_m3"]), num(hold["training_mean_baseline_overall"]["R2"]), num(hold["training_mean_baseline_overall"]["bias_pred_minus_actual_mg_m3"])),
        ],
        widths=[6.5, 2.2, 2.2, 1.8, 3.8],
        center_cols=(1, 2, 3, 4),
    )
    p(
        doc,
        "MAE и RMSE указаны в мг/м³; сдвиг — среднее (прогноз минус факт), также в мг/м³. "
        "На отложенном периоде итоговая модель заметно лучше простого среднего и обычной OLS. "
        "В 2015 году качество ниже, чем в 2014-м, а отрицательный сдвиг показывает "
        "недооценку концентрации. Это согласуется с более высоким средним CO в 2015 году."
    )
    p(
        doc,
        f"График показывает, что редкие очень высокие концентрации нередко недооцениваются. "
        f"На тесте получено {best['negative_predictions']} отрицательных прогноза из 14 542: "
        "физически они невозможны и требуют ограничения нулём в прикладной системе. "
        "Приведённые метрики рассчитаны до такого ограничения."
    )

    section_heading(doc, 6, "Выводы и возможности улучшения")
    p(
        doc,
        "Модель применима как вспомогательная оценка CO для изученной турбины и близких "
        "режимов работы. Она объясняет около 64 % вариации CO на данных последующих лет, "
        "но недостаточно надёжна для контроля редких пиков выбросов и не проверена на других "
        "установках. Нарушения условий гомоскедастичности и независимости ошибок ограничивают "
        "статистическую интерпретацию OLS; Ridge выбрана прежде всего ради прогноза."
    )
    for text in [
        "Собрать более свежие данные с точными временными метками и проверять модель по последовательным временным блокам.",
        "Добавить данные о топливе, ремонтах и изменениях режима, чтобы уменьшить влияние пропущенных факторов.",
        "Отдельно исследовать высокие концентрации: взвешенные потери, робастные методы и границы прогнозного интервала.",
        "Периодически переоценивать сдвиг прогноза по новым годам и проверить переносимость на других турбинах.",
    ]:
        list_item(doc, text)

    section_heading(doc, 7, "Список использованных источников и воспроизводимость")
    p(doc, "1. Gas Turbine CO and NOx Emission Data Set [Электронный ресурс] // UCI Machine Learning Repository. — 2019. — DOI: 10.24432/C5WC95. — URL: https://archive.ics.uci.edu/dataset/551 (дата обращения: 18.09.2026).")
    p(doc, "Исходный код: https://github.com/kotyasmol/gas-turbine-co-regression-lab. Основные файлы: analysis.py, make_report.py, requirements.txt и data/raw/gt_2011.csv … gt_2015.csv. Команда запуска приведена в README.md; численные результаты и диагностика — в outputs/results.json.")
    props = doc.core_properties
    author = " ".join(details["student"].split()[:2])
    props.author = author
    props.last_modified_by = author
    props.comments = ""
    props.title = "Лабораторная работа № 1 — линейная многомерная регрессия"
    props.subject = "Математические методы теории систем"
    props.keywords = ""
    now_utc = datetime.now(timezone.utc).replace(tzinfo=None)
    props.created = now_utc
    props.modified = now_utc
    doc.save(DEST)
    with ZipFile(DEST) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
        if "python-docx" in core.lower():
            raise RuntimeError("Generator metadata remains in docProps/core.xml")
    print(DEST)


if __name__ == "__main__":
    main()
